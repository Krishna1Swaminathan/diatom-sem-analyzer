import io

import openpyxl
import pytest

from diatom_analyzer.__main__ import main
from diatom_analyzer.classification import SpeciesLibrary, feature_distance
from diatom_analyzer.export import write_excel
from diatom_analyzer.pipeline import analyze_image
from diatom_analyzer.synthetic import DEMO_SCENE, render_scene
from tests.conftest import match_truth, png_bytes


@pytest.fixture()
def taught_library(tmp_path, demo_result):
    result, truth = demo_result
    lib = SpeciesLibrary(tmp_path / "lib")
    names = {"centric": "Big centric", "pennate": "Navicula", "girdle": "Girdle"}
    for gt, fr in match_truth(result, truth):
        if gt["damage"] == "intact" and gt["kind"] in names and gt["length_um"] >= 12:
            lib.add(names[gt["kind"]], fr, result.image, result.name, truth["um_per_px"])
    return SpeciesLibrary(tmp_path / "lib")


def test_library_round_trip(taught_library):
    assert set(taught_library.species) == {"Big centric", "Navicula", "Girdle"}


def test_species_identified_in_another_image(taught_library):
    img, truth = render_scene(DEMO_SCENE, seed=21, bar_style="none")
    res = analyze_image(png_bytes(img), name="other.png", manual_um_per_px=0.05, library=taught_library)
    got = {gt["kind"] + ("-big" if gt["length_um"] >= 12 else ""): fr for gt, fr in match_truth(res, truth)
           if gt["damage"] == "intact"}
    assert got["centric-big"].species == "Big centric"
    assert got["pennate-big"].species == "Navicula"
    assert got["girdle-big"].species == "Girdle"


def test_fragment_matched_by_pore_pattern(taught_library):
    img, truth = render_scene(DEMO_SCENE, seed=22, bar_style="none")
    res = analyze_image(png_bytes(img), name="frag.png", manual_um_per_px=0.05, library=taught_library)
    frag = next(fr for gt, fr in match_truth(res, truth) if gt["damage"] == "fragmented")
    assert frag.species == "Navicula"


def test_unfamiliar_shape_is_unknown(taught_library):
    from diatom_analyzer.synthetic import FrustuleSpec
    odd = [FrustuleSpec("pennate", 25, 19, 40, 3, angle_deg=10, pore_diameter_um=1.2, pore_spacing_um=2.4)]
    img, _ = render_scene(odd, bar_style="none", seed=23)
    res = analyze_image(png_bytes(img), name="odd.png", manual_um_per_px=0.05, library=taught_library)
    assert res.frustules[0].species == "unknown"


def test_missing_features_are_ignored():
    a = {"solidity": 0.95, "log_length_um": float("nan")}
    b = {"solidity": 0.95, "log_length_um": 2.0}
    assert feature_distance(a, b) == 0.0


def test_excel_workbook(demo_result):
    result, _ = demo_result
    buf = io.BytesIO()
    write_excel([result], buf)
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Summary", "Frustules", "Pores", "Settings", "Column guide"]
    frustules = wb["Frustules"]
    header = [c.value for c in frustules[1]]
    for col in ("frustule_id", "damage", "morphotype", "length_um", "orientation_deg", "pore_count"):
        assert col in header
    assert frustules.max_row == 1 + len(result.frustules)
    pores = wb["Pores"]
    pore_header = [c.value for c in pores[1]]
    assert {"x_um", "y_um", "diameter_um"} <= set(pore_header)
    assert pores.max_row == 1 + sum(len(f.pores) for f in result.frustules)


def test_cli_batch(tmp_path):
    for seed in (1, 2):
        img, _ = render_scene(DEMO_SCENE[:3], seed=seed, bar_style="none")
        (tmp_path / f"img{seed}.png").write_bytes(png_bytes(img))
    out = tmp_path / "out.xlsx"
    code = main([str(tmp_path), "-o", str(out), "--um-per-px", "0.05", "--overlays", str(tmp_path / "ov"),
                 "--library", str(tmp_path / "nolib")])
    assert code == 0
    wb = openpyxl.load_workbook(out)
    assert wb["Summary"].max_row == 3
    assert len(list((tmp_path / "ov").glob("*.png"))) == 2
